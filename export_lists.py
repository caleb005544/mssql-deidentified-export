# generate_reports.py
# Timeseries report generator (時序列速報) — clean v1.0
# - Supports scope: brand / sv / store
# - Auto month selection if year_month is None (based on MAX(date) in measure table under the scope)
# - Output format:
#   A: 日期 (yyyymmdd 星期X)
#   Then for each store: 日商 / 來客數 / 客單價
# - Excel output uses openpyxl (handles MultiIndex columns cleanly)
#
# Prereq:
#   pip install -r requirements.txt
#
# Run:
#   python generate_reports.py
#
# Notes:
#   - This file assumes Postgres is reachable at localhost:5433, db reporting/reporting.
#   - "brand" values in your data look like: "brand 1", "brand 2", "brand 3"

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import List, Optional, Tuple

import pandas as pd
from dotenv import load_dotenv
from sqlalchemy import create_engine, text
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# -----------------------------
# Config
# -----------------------------

load_dotenv()

DB_HOST = os.getenv("PGHOST", "localhost")
DB_PORT = int(os.getenv("PGPORT", "5433"))
DB_NAME = os.getenv("PGDATABASE", "reporting")
DB_USER = os.getenv("PGUSER", "reporting")
DB_PASS = os.getenv("PGPASSWORD", "reporting")

OUTPUT_DIR = os.getenv("OUTPUT_DIR", "output")

# IMPORTANT: you renamed calc -> measure
MEASURE_TABLE = os.getenv("MEASURE_TABLE", "measure.daily_store_summary")

# Use psycopg v3 driver (installed by requirements.txt: psycopg)
# SQLAlchemy 2.x: postgresql+psycopg://
ENGINE = create_engine(
    f"postgresql+psycopg://{DB_USER}:{DB_PASS}@{DB_HOST}:{DB_PORT}/{DB_NAME}",
    future=True,
)

METRICS_ORDER = ["日商", "來客數", "客單價"]
METRIC_MAP = {"revenue": "日商", "customer_count": "來客數", "average_ticket": "客單價"}

# -----------------------------
# Utilities
# -----------------------------


def ensure_output_dir() -> None:
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def safe_filename(s: str) -> str:
    s = s.strip()
    s = re.sub(r"\s+", "_", s)
    s = re.sub(r"[^0-9A-Za-z_\-]+", "", s)
    return s or "value"


def month_date_range(year_month: str) -> Tuple[str, str]:
    """
    year_month: 'YYYY-MM'
    return: (start_date, end_date) as 'YYYY-MM-DD'
    """
    y, m = year_month.split("-")
    y = int(y)
    m = int(m)
    start = date(y, m, 1)
    if m == 12:
        end = date(y + 1, 1, 1) - pd.Timedelta(days=1)
    else:
        end = date(y, m + 1, 1) - pd.Timedelta(days=1)
    return start.isoformat(), end.isoformat()


def resolve_year_month(scope_type: str, scope_value: str, year_month: Optional[str]) -> str:
    """
    If year_month is provided: return it.
    Else: infer from MAX(date) within the store scope.
    """
    if year_month:
        return year_month

    stores = get_store_scope(scope_type, scope_value)
    if not stores:
        raise ValueError(f"No stores found for {scope_type}={scope_value}")

    sql = text(
        f"""
        SELECT to_char(MAX(date), 'YYYY-MM') AS ym
        FROM {MEASURE_TABLE}
        WHERE storeno = ANY(:stores)
        """
    )
    with ENGINE.begin() as conn:
        ym = conn.execute(sql, {"stores": stores}).scalar()

    if not ym:
        # fallback: current month
        return datetime.today().strftime("%Y-%m")

    return str(ym)


# -----------------------------
# Store scope resolution
# -----------------------------


def get_store_scope(scope_type: str, scope_value: str) -> List[str]:
    """
    scope_type: 'brand' | 'sv' | 'store'
    - brand: semantic.dim_store.brand = scope_value => storeno list
    - sv: semantic.dim_sv.sv_name = scope_value => storeno list
    - store: semantic.dim_store.store_name = scope_value OR storeno = scope_value => single
    """
    scope_type = scope_type.strip().lower()

    if scope_type == "brand":
        sql = text(
            """
            SELECT DISTINCT storeno
            FROM semantic.dim_store
            WHERE brand = :brand
            ORDER BY storeno;
            """
        )
        with ENGINE.begin() as conn:
            rows = conn.execute(sql, {"brand": scope_value}).fetchall()
        return [r[0] for r in rows]

    if scope_type == "sv":
        sql = text(
            """
            SELECT DISTINCT storeno
            FROM semantic.dim_sv
            WHERE sv_name = :sv
            ORDER BY storeno;
            """
        )
        with ENGINE.begin() as conn:
            rows = conn.execute(sql, {"sv": scope_value}).fetchall()
        return [r[0] for r in rows]

    if scope_type == "store":
        # allow passing storeno directly (e.g., B1002) OR store_name (e.g., Store 1002)
        sql = text(
            """
            SELECT storeno
            FROM semantic.dim_store
            WHERE store_name = :store_name
               OR storeno = :storeno
            ORDER BY storeno
            LIMIT 1;
            """
        )
        with ENGINE.begin() as conn:
            storeno = conn.execute(sql, {"store_name": scope_value, "storeno": scope_value}).scalar()
        return [storeno] if storeno else []

    raise ValueError("scope_type must be one of: brand | sv | store")


# -----------------------------
# Excel writer
# -----------------------------


def write_timeseries_excel(pivot: pd.DataFrame, date_labels: pd.Series, out_path: str) -> None:
    """
    pivot:
      - columns are MultiIndex (store_name, metric)
      - values are numeric (already rounded)
    date_labels:
      - Series of 'yyyymmdd 星期X'
    Output:
      A = 日期
      For each store: 日商/來客數/客單價 (3 cols)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "時序列速報"

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_font = Font(bold=True)

    # Discover stores from MultiIndex columns
    mi_cols = [c for c in pivot.columns if isinstance(c, tuple) and len(c) == 2]
    stores = []
    for s, _m in mi_cols:
        if s not in stores:
            stores.append(s)

    # Header rows
    ws.cell(row=1, column=1, value="日期").alignment = center
    ws.cell(row=1, column=1).font = header_font
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    col_ptr = 2
    for s in stores:
        ws.cell(row=1, column=col_ptr, value=str(s)).alignment = center
        ws.cell(row=1, column=col_ptr).font = header_font
        ws.merge_cells(start_row=1, start_column=col_ptr, end_row=1, end_column=col_ptr + 2)

        for i, m in enumerate(METRICS_ORDER):
            ws.cell(row=2, column=col_ptr + i, value=m).alignment = center
            ws.cell(row=2, column=col_ptr + i).font = header_font

        col_ptr += 3

    # Body
    start_row = 3
    for i in range(len(date_labels)):
        r = start_row + i
        ws.cell(row=r, column=1, value=str(date_labels.iloc[i]))

        col_ptr = 2
        for s in stores:
            for m in METRICS_ORDER:
                val = None
                if (s, m) in pivot.columns:
                    val = pivot.iloc[i][(s, m)]
                c = ws.cell(row=r, column=col_ptr, value=None if pd.isna(val) else float(val))
                c.number_format = "#,##0"  # thousands + integer
                col_ptr += 1

    # Column widths
    ws.column_dimensions["A"].width = 16
    total_cols = 1 + len(stores) * 3
    for c in range(2, total_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10

    wb.save(out_path)


# -----------------------------
# Report generator
# -----------------------------

def generate_timeseries_report(scope_type: str, scope_value: str, year_month: Optional[str] = None) -> Optional[str]:
    ensure_output_dir()

    stores = get_store_scope(scope_type, scope_value)
    if not stores:
        print(f"[WARN] No stores found for {scope_type}={scope_value}")
        return None

    ym = resolve_year_month(scope_type, scope_value, year_month)
    start_date, end_date = month_date_range(ym)

    sql = text(
        f"""
        WITH base AS (
            SELECT
                date,
                storeno,
                store_name,
                SUM(revenue)::numeric AS revenue,
                SUM(customer_count)::bigint AS customer_count
            FROM {MEASURE_TABLE}
            WHERE date BETWEEN :start_date AND :end_date
              AND storeno = ANY(:stores)
            GROUP BY date, storeno, store_name
        )
        SELECT
            b.date,
            b.store_name,
            b.revenue,
            b.customer_count,
            CASE WHEN b.customer_count = 0 THEN NULL ELSE b.revenue / b.customer_count END AS average_ticket
        FROM base b
        ORDER BY b.date, b.store_name;
        """
    )

    with ENGINE.begin() as conn:
        df = pd.read_sql(sql, conn, params={"start_date": start_date, "end_date": end_date, "stores": stores})

    if df.empty:
        print(f"[WARN] No data for {scope_type}={scope_value} in {ym}")
        return None

    weekday_sql = text(
        """
        SELECT date, weekday_name
        FROM semantic.dim_date
        WHERE date BETWEEN :start_date AND :end_date
        ORDER BY date;
        """
    )
    with ENGINE.begin() as conn:
        wd = pd.read_sql(weekday_sql, conn, params={"start_date": start_date, "end_date": end_date})

    # ✅ critical: ensure datetime-like BEFORE using .dt / merging
    df["date"] = pd.to_datetime(df["date"])
    wd["date"] = pd.to_datetime(wd["date"])

    df = df.merge(wd, on="date", how="left")

    df["yyyymmdd"] = df["date"].dt.strftime("%Y%m%d")
    df["date_label"] = df["yyyymmdd"] + " " + df["weekday_name"].fillna("")

    # Round to integer
    df["revenue"] = pd.to_numeric(df["revenue"], errors="coerce").round(0)
    df["customer_count"] = pd.to_numeric(df["customer_count"], errors="coerce").round(0)
    df["average_ticket"] = pd.to_numeric(df["average_ticket"], errors="coerce").round(0)

    long = df.melt(
        id_vars=["date", "date_label", "store_name"],
        value_vars=["revenue", "customer_count", "average_ticket"],
        var_name="metric",
        value_name="value",
    )

    metric_map = {"revenue": "日商", "customer_count": "來客數", "average_ticket": "客單價"}
    long["metric"] = long["metric"].map(metric_map)
    long["metric"] = pd.Categorical(long["metric"], categories=["日商", "來客數", "客單價"], ordered=True)

    pivot = (
        long.pivot_table(
            index=["date", "date_label"],
            columns=["store_name", "metric"],
            values="value",
            aggfunc="sum",
        )
        .reset_index()
        .sort_values("date")
        .reset_index(drop=True)
    )

    # ✅ only keep date_label as the first column in Excel (remove date_label extra column issue)
    date_labels = pivot["date_label"].copy()
    pivot = pivot.drop(columns=["date", "date_label"])

    scope_safe = safe_filename(scope_value)
    out_name = f"時序列速報_{scope_type}_{scope_safe}_{ym}.xlsx"
    out_path = os.path.join(OUTPUT_DIR, out_name)

    write_timeseries_excel(pivot, date_labels, out_path)

    print(f"[OK] Report generated: {out_path}")
    return out_path




# -----------------------------
# Main
# -----------------------------

if __name__ == "__main__":
    # Example:
    #   brand values: "brand 1" / "brand 2" / "brand 3"
    #   year_month can be None to auto-pick from max(date) under scope
    generate_timeseries_report("brand", "brand 1", year_month=None)
    # generate_timeseries_report("store", "Store 1002", year_month="2025-12")
    # generate_timeseries_report("sv", "SV_01", year_month=None)
